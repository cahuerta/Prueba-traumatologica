"""
routers/alumnos.py
Carga y gestión del listado de alumnos del conjunto (generación) activo.
Dos formas de carga, ambas contra el conjunto activo, ambas por upsert
sobre (conjunto_id, rut) -no porque exista una lista previa dentro del
mismo conjunto, sino por robustez: evita duplicados si se sube el Excel
dos veces, colapsa RUTs repetidos dentro del mismo archivo, y permite
agregar un alumno nuevo a mano después de la carga inicial sin afectar
a los demás-.

Solo el admin gestiona el listado de alumnos: es quien decide qué
generación/conjunto está en uso y quiénes la componen.
"""

import io
from typing import List

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from pydantic import BaseModel
from openpyxl import load_workbook

from routers.auth import sb, get_current_interrogador, requiere_admin
from routers.conjuntos_comun import obtener_conjunto_activo_id

router = APIRouter(prefix="/alumnos", tags=["alumnos"])


# ---------------- MODELOS ----------------
class AlumnoIn(BaseModel):
    rut: str
    nombre: str

class CargaAlumnosIn(BaseModel):
    alumnos: List[AlumnoIn]


# ---------------- HELPER COMPARTIDO ----------------
def _upsert_alumnos(alumnos: List[AlumnoIn], conjunto_id: str) -> dict:
    creados = 0
    actualizados = 0

    for a in alumnos:
        rut = a.rut.strip()
        nombre = a.nombre.strip()
        if not rut or not nombre:
            continue

        existente = sb.table("alumnos").select("id").eq("rut", rut).eq("conjunto_id", conjunto_id).execute().data
        if existente:
            sb.table("alumnos").update({"nombre": nombre}).eq("id", existente[0]["id"]).execute()
            actualizados += 1
        else:
            sb.table("alumnos").insert({"rut": rut, "nombre": nombre, "conjunto_id": conjunto_id}).execute()
            creados += 1

    return {"creados": creados, "actualizados": actualizados}


# ---------------- ENDPOINTS ----------------
@router.get("")
def listar_alumnos(interrogador: dict = Depends(get_current_interrogador)):
    """Lista los alumnos del conjunto actualmente activo."""
    conjunto_id = obtener_conjunto_activo_id()
    return sb.table("alumnos").select("*").eq("conjunto_id", conjunto_id).order("nombre").execute().data


@router.post("")
def cargar_alumnos_json(body: CargaAlumnosIn, admin: dict = Depends(requiere_admin)):
    """Carga (o agrega) alumnos al conjunto activo, pegados como lista
    ya parseada (nombre + RUT)."""
    conjunto_id = obtener_conjunto_activo_id()
    resultado = _upsert_alumnos(body.alumnos, conjunto_id)
    return {"conjunto_id": conjunto_id, **resultado}


@router.post("/excel")
async def cargar_alumnos_excel(
    archivo: UploadFile = File(...),
    admin: dict = Depends(requiere_admin),
):
    """Carga alumnos al conjunto activo desde un archivo .xlsx.
    Primera fila es encabezado: RUT, Nombre, Apellido (en ese orden).
    Nombre y Apellido se concatenan en un solo campo 'nombre'."""
    conjunto_id = obtener_conjunto_activo_id()

    contenido = await archivo.read()
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro.active
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo Excel")

    filas = list(hoja.iter_rows(min_row=2, values_only=True))  # se salta el encabezado

    alumnos = []
    for fila in filas:
        if not fila or len(fila) < 2:
            continue
        rut = str(fila[0]).strip() if fila[0] else ""
        nombre_raw = str(fila[1]).strip() if len(fila) > 1 and fila[1] else ""
        apellido_raw = str(fila[2]).strip() if len(fila) > 2 and fila[2] else ""
        nombre_completo = f"{nombre_raw} {apellido_raw}".strip()

        if rut and nombre_completo:
            alumnos.append(AlumnoIn(rut=rut, nombre=nombre_completo))

    if not alumnos:
        raise HTTPException(422, "El archivo no contiene filas válidas (RUT + Nombre)")

    resultado = _upsert_alumnos(alumnos, conjunto_id)
    return {"conjunto_id": conjunto_id, "total_filas_leidas": len(filas), **resultado}


@router.delete("/{alumno_id}")
def borrar_alumno(alumno_id: str, admin: dict = Depends(requiere_admin)):
    sb.table("alumnos").delete().eq("id", alumno_id).execute()
    return {"ok": True}
      
