"""
services/avatar_clase.py
Prepara la respuesta hablada del avatar para una pregunta de alumno en
Clases Formales. Se ejecuta solo cuando el interrogador toca "Preparar".

Flujo (para no gastar de mas):
1. Filtro sin costo: frases muy cortas no se procesan.
2. Clasificacion (modelo barato, sin materiales): ¿es de traumatologia?,
   ¿de que region?, y terminos de busqueda clinicos para la literatura.
3. Material del curso de esa region (.docx y .pptx, con cache en memoria):
   respuesta corta con el modelo completo, SOLO con lo que dice el material.
4. Respaldo inmediato en literatura (EvidenciaMed, server-to-server) si el
   material no alcanza o la region no tiene material: /search + /analyze/doi
   con 1 paper, resumido para voz y aclarando que no viene del curso.

Reutiliza la conexion a Supabase y el lector de .docx de fundamento_vivo.py
sin modificarlo.

Variables de entorno:
  ANTHROPIC_API_KEY
  EVIDENCIAMED_URL        URL del SERVIDOR de EvidenciaMed (sin "/" final)
  EVIDENCIAMED_API_KEY    la misma X-API-Key que usa su frontend
  AVATAR_MODELO_CLASIFICADOR  (opcional, por defecto claude-haiku-4-5)
  AVATAR_MODELO_RESPUESTA     (opcional, por defecto claude-sonnet-4-6)
"""

import io
import json
import os
import re
import threading
import time
import unicodedata
from typing import Dict, List, Optional, Tuple

import anthropic
import httpx
from pptx import Presentation

from routers.auth import sb
from services.fundamento_vivo import _extraer_texto_docx

client = anthropic.Anthropic()  # usa ANTHROPIC_API_KEY del entorno

MODELO_CLASIFICADOR = os.getenv("AVATAR_MODELO_CLASIFICADOR", "claude-haiku-4-5")
MODELO_RESPUESTA = os.getenv("AVATAR_MODELO_RESPUESTA", "claude-sonnet-4-6")

EVIDENCIAMED_URL = os.getenv("EVIDENCIAMED_URL", "").rstrip("/")
EVIDENCIAMED_API_KEY = os.getenv("EVIDENCIAMED_API_KEY", "")

MAX_PALABRAS = 70            # limite de la respuesta hablada
MIN_PALABRAS_PREGUNTA = 3
CACHE_TTL_SEG = 15 * 60
MAX_CARACTERES_REGION = 150_000
NO_ESTA = "NO_ESTA_EN_MATERIAL"

# ---------------- utilidades ----------------
def _normalizar(texto: str) -> str:
    s = unicodedata.normalize("NFD", texto or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.lower()).strip()


def _texto_de(mensaje) -> str:
    return "".join(b.text for b in mensaje.content if getattr(b, "type", "") == "text").strip()


def _json_de(texto: str) -> dict:
    limpio = texto.replace("```json", "").replace("```", "").strip()
    m = re.search(r"\{.*\}", limpio, re.DOTALL)
    return json.loads(m.group(0) if m else limpio)


def recortar_palabras(texto: str, maximo: int = MAX_PALABRAS) -> str:
    """Asegura el limite de palabras, cortando en el ultimo punto posible."""
    palabras = (texto or "").split()
    if len(palabras) <= maximo:
        return " ".join(palabras)
    corto = " ".join(palabras[:maximo])
    ultimo_punto = max(corto.rfind("."), corto.rfind("?"), corto.rfind("!"))
    if ultimo_punto > len(corto) * 0.5:
        return corto[: ultimo_punto + 1]
    return corto.rstrip(",;:") + "."


# ---------------- material del curso (.docx y .pptx) ----------------
def _extraer_texto_pptx(contenido: bytes) -> str:
    """Texto de todas las diapositivas (cuadros de texto y tablas) + notas del orador."""
    presentacion = Presentation(io.BytesIO(contenido))
    lineas: List[str] = []
    for numero, diapositiva in enumerate(presentacion.slides, start=1):
        partes: List[str] = []
        for forma in diapositiva.shapes:
            if getattr(forma, "has_text_frame", False) and forma.has_text_frame:
                texto = forma.text_frame.text.strip()
                if texto:
                    partes.append(texto)
            if getattr(forma, "has_table", False) and forma.has_table:
                for fila in forma.table.rows:
                    celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                    if celdas:
                        partes.append(" | ".join(celdas))
        if diapositiva.has_notes_slide:
            notas = diapositiva.notes_slide.notes_text_frame.text.strip()
            if notas:
                partes.append(f"Notas: {notas}")
        if partes:
            lineas.append(f"[Diapositiva {numero}] " + "\n".join(partes))
    return "\n".join(lineas)


_lock_cache = threading.Lock()
_cache_regiones: Dict[str, object] = {"datos": None, "ts": 0.0}
_cache_texto: Dict[str, Tuple[str, float]] = {}


def _filas_legibles() -> List[dict]:
    filas = sb.table("materiales").select("id, titulo, storage_path, region").execute().data or []
    return [
        f for f in filas
        if f.get("region") and (f.get("storage_path") or "").lower().endswith((".docx", ".pptx"))
    ]


def listar_regiones() -> List[str]:
    if _cache_regiones["datos"] is not None and time.time() - _cache_regiones["ts"] < CACHE_TTL_SEG:
        return list(_cache_regiones["datos"])
    regiones = sorted({f["region"] for f in _filas_legibles()})
    _cache_regiones["datos"], _cache_regiones["ts"] = regiones, time.time()
    return list(regiones)


def texto_region(region: str) -> str:
    guardado = _cache_texto.get(region)
    if guardado and time.time() - guardado[1] < CACHE_TTL_SEG:
        return guardado[0]
    with _lock_cache:
        guardado = _cache_texto.get(region)
        if guardado and time.time() - guardado[1] < CACHE_TTL_SEG:
            return guardado[0]
        bloques: List[str] = []
        for f in [f for f in _filas_legibles() if f["region"] == region]:
            try:
                contenido = sb.storage.from_("materiales").download(f["storage_path"])
                if f["storage_path"].lower().endswith(".pptx"):
                    texto = _extraer_texto_pptx(contenido)
                else:
                    texto = _extraer_texto_docx(contenido)
            except Exception as e:  # noqa: BLE001 — material ilegible: se omite
                print(f"[avatar_clase] no se pudo leer {f['storage_path']}: {e!r}")
                continue
            if texto.strip():
                bloques.append(f"### {f['titulo']}\n{texto}")
        total = "\n\n".join(bloques)[:MAX_CARACTERES_REGION]
        _cache_texto[region] = (total, time.time())
        return total


# ---------------- pasos con Claude ----------------
def clasificar(pregunta: str, regiones: List[str]) -> dict:
    prompt = (
        "Clasifica la pregunta de un alumno de 4to año de medicina en una clase de traumatología.\n\n"
        f"PREGUNTA: {pregunta}\n\n"
        f"REGIONES CON MATERIAL (usa exactamente uno de estos textos): {json.dumps(regiones, ensure_ascii=False)}\n\n"
        "Reglas:\n"
        "- pertinente = true solo si es sobre traumatología u ortopedia (lesiones, fracturas, patología "
        "musculoesquelética, examen físico, imagenología, tratamiento).\n"
        "- region = la región con material a la que corresponde, o null si no calza con ninguna.\n"
        "- terminos = 3 a 6 palabras clínicas en español para buscar literatura sobre la pregunta.\n\n"
        'Responde SOLO con JSON válido: {"pertinente": true|false, "region": "texto"|null, "terminos": "..."}'
    )
    datos = _json_de(_texto_de(client.messages.create(
        model=MODELO_CLASIFICADOR, max_tokens=150, messages=[{"role": "user", "content": prompt}],
    )))
    region = None
    if isinstance(datos.get("region"), str):
        objetivo = _normalizar(datos["region"])
        region = next((r for r in regiones if _normalizar(r) == objetivo), None)
    return {
        "pertinente": bool(datos.get("pertinente")),
        "region": region,
        "terminos": str(datos.get("terminos") or pregunta).strip(),
    }


def responder_con_material(pregunta: str, region: str, contexto: str) -> Optional[str]:
    """Devuelve la respuesta, o None si el material no permite responder."""
    prompt = (
        "Eres una docente de traumatología que responde EN VOZ ALTA, frente al curso, la pregunta "
        "de un alumno de 4to año de medicina. Una voz sintética leerá tu respuesta.\n\n"
        f"REGIÓN: {region}\nPREGUNTA: {pregunta}\n\n"
        "MATERIAL DEL CURSO (documentos y presentaciones de los docentes, cada uno bajo '### Título'):\n"
        f"{contexto}\n\n"
        "Instrucciones:\n"
        "- Responde SOLO con lo que está en el material. No agregues información externa.\n"
        f"- Si el material no permite responder, escribe exactamente {NO_ESTA} y nada más.\n"
        f"- Máximo {MAX_PALABRAS} palabras, en 2 a 4 frases claras, español de Chile.\n"
        "- Sin listas, viñetas, títulos, asteriscos ni formato: solo texto corrido para ser hablado.\n"
        "- No menciones los títulos de los documentos ni las diapositivas."
    )
    texto = _texto_de(client.messages.create(
        model=MODELO_RESPUESTA, max_tokens=300, messages=[{"role": "user", "content": prompt}],
    ))
    if not texto or NO_ESTA in texto:
        return None
    return recortar_palabras(texto)


def _primer_autor(autores: str) -> str:
    primero = (autores or "").split(",")[0].strip()
    return primero.split()[0] if primero else ""


def responder_con_literatura(pregunta: str, terminos: str) -> Optional[str]:
    """1 paper de EvidenciaMed, resumido para voz. None si no hay resultado."""
    if not EVIDENCIAMED_URL or not EVIDENCIAMED_API_KEY:
        print("[avatar_clase] EvidenciaMed no configurado (EVIDENCIAMED_URL / EVIDENCIAMED_API_KEY)")
        return None
    cabeceras = {"X-API-Key": EVIDENCIAMED_API_KEY}
    try:
        with httpx.Client(timeout=httpx.Timeout(150.0, connect=60.0)) as c:
            r = c.post(f"{EVIDENCIAMED_URL}/search", headers=cabeceras,
                       json={"query": terminos, "max_results": 5})
            if r.status_code != 200:
                print(f"[avatar_clase] EvidenciaMed /search HTTP {r.status_code}")
                return None
            papers = [p for p in (r.json().get("papers") or []) if p.get("doi")]
            if not papers:
                print(f"[avatar_clase] EvidenciaMed sin papers con DOI para: {terminos}")
                return None
            paper = papers[0]
            r = c.post(f"{EVIDENCIAMED_URL}/analyze/doi", headers=cabeceras, json={"doi": paper["doi"]})
            if r.status_code != 200:
                print(f"[avatar_clase] EvidenciaMed /analyze/doi HTTP {r.status_code}")
                return None
            analisis = r.json()
    except Exception as e:  # noqa: BLE001
        print(f"[avatar_clase] EvidenciaMed error de conexión: {e!r}")
        return None

    autor = _primer_autor(analisis.get("autores") or paper.get("authors", ""))
    anio = paper.get("year") or ""
    resumen = {
        "titulo": analisis.get("titulo") or paper.get("title"),
        "tipo_estudio": analisis.get("tipo_estudio"),
        "nivel_evidencia": analisis.get("nivel_evidencia_oxford"),
        "resumen": analisis.get("resumen_ejecutivo"),
        "hallazgos": analisis.get("hallazgos_clave"),
    }
    prompt = (
        "Eres una docente de traumatología que responde EN VOZ ALTA la pregunta de un alumno. "
        "El tema no está en el material del curso, así que respondes con un estudio científico.\n\n"
        f"PREGUNTA: {pregunta}\n\n"
        f"ESTUDIO (autor principal: {autor or 'desconocido'}, año: {anio or 'desconocido'}):\n"
        f"{json.dumps(resumen, ensure_ascii=False)}\n\n"
        "Instrucciones:\n"
        "- Empieza diciendo que no está en el material del curso y que según un estudio "
        "de ese autor y año, y responde con lo que dice el estudio.\n"
        "- Si el estudio no responde la pregunta, dilo honestamente en una frase.\n"
        f"- Máximo {MAX_PALABRAS} palabras, español de Chile, sin listas ni formato."
    )
    texto = _texto_de(client.messages.create(
        model=MODELO_CLASIFICADOR, max_tokens=300, messages=[{"role": "user", "content": prompt}],
    ))
    return recortar_palabras(texto) if texto else None


# ---------------- orquestacion ----------------
def preparar_respuesta(pregunta: str) -> dict:
    """
    Devuelve {estado, respuesta, fuente, region}.
    estado: lista | no_pertinente | error
    """
    pregunta = (pregunta or "").strip()
    if len(pregunta.split()) < MIN_PALABRAS_PREGUNTA:
        return {"estado": "no_pertinente", "respuesta": "", "fuente": "ninguna", "region": None}

    try:
        regiones = listar_regiones()
    except Exception as e:  # noqa: BLE001
        print(f"[avatar_clase] no se pudieron listar regiones: {e!r}")
        regiones = []

    try:
        clase = clasificar(pregunta, regiones)
    except Exception as e:  # noqa: BLE001
        print(f"[avatar_clase] error clasificando: {e!r}")
        return {"estado": "error", "respuesta": "", "fuente": "ninguna", "region": None}

    if not clase["pertinente"]:
        return {"estado": "no_pertinente", "respuesta": "", "fuente": "ninguna", "region": None}

    region = clase["region"]
    if region:
        try:
            contexto = texto_region(region)
            if contexto.strip():
                respuesta = responder_con_material(pregunta, region, contexto)
                if respuesta:
                    return {"estado": "lista", "respuesta": respuesta, "fuente": "curso", "region": region}
        except Exception as e:  # noqa: BLE001
            print(f"[avatar_clase] error con material de {region}: {e!r}")

    respuesta = responder_con_literatura(pregunta, clase["terminos"])
    if respuesta:
        return {"estado": "lista", "respuesta": respuesta, "fuente": "literatura", "region": region}

    return {"estado": "error", "respuesta": "", "fuente": "ninguna", "region": region}
